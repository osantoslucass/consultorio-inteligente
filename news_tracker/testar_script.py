#!/usr/bin/env python3
"""Testes unitários com dados mockados — valida lógica de Excel e deduplicação."""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(__file__))

from buscar_noticias import carregar_links_existentes, salvar_noticias
import openpyxl

NOTICIAS_MOCK = [
    {
        "titulo": "Plano de saúde amplia cobertura odontológica",
        "link": "https://exemplo.com/noticias/1",
        "data": "2026-06-01 08:00",
        "resumo": "Novos convênios incluem tratamentos de canal e implantes.",
    },
    {
        "titulo": "Dentistas adotam IA para diagnóstico de cáries",
        "link": "https://exemplo.com/noticias/2",
        "data": "2026-06-01 09:00",
        "resumo": "Tecnologia reduz tempo de diagnóstico em 40%.",
    },
    {
        "titulo": "Saúde bucal e doenças cardiovasculares: nova pesquisa",
        "link": "https://exemplo.com/noticias/3",
        "data": "2026-06-01 10:00",
        "resumo": "Estudo da USP relaciona periodontite a inflamações cardíacas.",
    },
]


class TestSalvarNoticias(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mktemp(suffix=".xlsx")

    def tearDown(self):
        if os.path.exists(self.tmp):
            os.remove(self.tmp)

    def test_cria_arquivo_novo(self):
        inseridas = salvar_noticias(self.tmp, NOTICIAS_MOCK, "odontologia", set())
        self.assertEqual(inseridas, 3)
        self.assertTrue(os.path.exists(self.tmp))

    def test_cabecalho_correto(self):
        salvar_noticias(self.tmp, NOTICIAS_MOCK, "odontologia", set())
        wb = openpyxl.load_workbook(self.tmp)
        ws = wb.active
        cabecalho = [cell.value for cell in ws[1]]
        self.assertIn("Título", cabecalho)
        self.assertIn("Link", cabecalho)
        self.assertIn("Resumo", cabecalho)

    def test_deduplica_em_segunda_execucao(self):
        links = set()
        salvar_noticias(self.tmp, NOTICIAS_MOCK, "odontologia", links)

        links_apos = carregar_links_existentes(self.tmp)
        inseridas2 = salvar_noticias(self.tmp, NOTICIAS_MOCK, "odontologia", links_apos)
        self.assertEqual(inseridas2, 0, "Não deve inserir duplicatas")

    def test_adiciona_apenas_novos(self):
        links = set()
        salvar_noticias(self.tmp, NOTICIAS_MOCK[:2], "odontologia", links)

        links_apos = carregar_links_existentes(self.tmp)
        # Envia todas de novo — só a 3ª é nova
        inseridas2 = salvar_noticias(self.tmp, NOTICIAS_MOCK, "odontologia", links_apos)
        self.assertEqual(inseridas2, 1)

    def test_carregar_links_arquivo_inexistente(self):
        links = carregar_links_existentes("/nao/existe.xlsx")
        self.assertIsInstance(links, set)
        self.assertEqual(len(links), 0)

    def test_conteudo_celulas(self):
        salvar_noticias(self.tmp, [NOTICIAS_MOCK[0]], "odontologia", set())
        wb = openpyxl.load_workbook(self.tmp)
        ws = wb.active
        linha2 = [cell.value for cell in ws[2]]
        self.assertEqual(linha2[0], NOTICIAS_MOCK[0]["titulo"])
        self.assertEqual(linha2[1], NOTICIAS_MOCK[0]["link"])
        self.assertEqual(linha2[5], "odontologia")


if __name__ == "__main__":
    unittest.main(verbosity=2)
