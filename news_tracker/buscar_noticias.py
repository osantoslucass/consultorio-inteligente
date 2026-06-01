#!/usr/bin/env python3
"""
Rotina de acompanhamento de notícias.
Busca notícias via Google News RSS, salva em planilha.xlsx sem duplicar.
"""

import os
import sys
import argparse
import urllib.request
import urllib.parse
import xml.etree.ElementTree as ET
from datetime import datetime
from email.utils import parsedate_to_datetime
from bs4 import BeautifulSoup
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from dotenv import load_dotenv

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))

COLUNAS = ["Título", "Link", "Data Publicação", "Resumo", "Data de Coleta", "Tema"]
CABECALHO_COR = "1F4E79"
CABECALHO_FONTE_COR = "FFFFFF"


def parse_args():
    parser = argparse.ArgumentParser(description="Busca notícias e salva em Excel")
    parser.add_argument(
        "--tema",
        default=os.getenv("TEMA_NOTICIAS", "odontologia dentista"),
        help="Tema de busca (padrão: variável TEMA_NOTICIAS ou 'odontologia dentista')",
    )
    parser.add_argument(
        "--arquivo",
        default=os.getenv("ARQUIVO_EXCEL", "planilha.xlsx"),
        help="Caminho do arquivo Excel (padrão: planilha.xlsx)",
    )
    parser.add_argument(
        "--max",
        type=int,
        default=30,
        help="Número máximo de notícias a buscar (padrão: 30)",
    )
    return parser.parse_args()


def buscar_noticias_google(tema: str, max_resultados: int) -> list[dict]:
    """Busca notícias no Google News RSS em português (Brasil)."""
    query = urllib.parse.quote(tema)
    url = (
        f"https://news.google.com/rss/search"
        f"?q={query}&hl=pt-BR&gl=BR&ceid=BR:pt-419"
    )

    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            conteudo = resp.read()
    except Exception as e:
        print(f"[ERRO] Falha ao acessar Google News: {e}", file=sys.stderr)
        return []

    try:
        root = ET.fromstring(conteudo)
    except ET.ParseError as e:
        print(f"[ERRO] Falha ao parsear RSS: {e}", file=sys.stderr)
        return []

    ns = {"media": "http://search.yahoo.com/mrss/"}
    noticias = []

    for item in root.findall(".//item")[:max_resultados]:
        titulo_raw = item.findtext("title", "").strip()
        link = item.findtext("link", "").strip()
        data_str = item.findtext("pubDate", "").strip()
        descricao_raw = item.findtext("description", "").strip()

        # Limpa título — Google News inclui " - Fonte" no final
        titulo = titulo_raw.rsplit(" - ", 1)[0] if " - " in titulo_raw else titulo_raw

        # Normaliza data
        try:
            data_pub = parsedate_to_datetime(data_str).strftime("%Y-%m-%d %H:%M")
        except Exception:
            data_pub = data_str

        # Remove HTML do resumo
        resumo = BeautifulSoup(descricao_raw, "html.parser").get_text(separator=" ").strip()
        resumo = " ".join(resumo.split())[:400]

        if link:
            noticias.append(
                {
                    "titulo": titulo,
                    "link": link,
                    "data": data_pub,
                    "resumo": resumo,
                }
            )

    return noticias


def carregar_links_existentes(arquivo: str) -> set:
    """Retorna o conjunto de links já presentes na planilha."""
    if not os.path.exists(arquivo):
        return set()
    try:
        wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
        ws = wb.active
        links = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[1]:
                links.add(str(row[1]).strip())
        wb.close()
        return links
    except Exception as e:
        print(f"[AVISO] Não foi possível ler planilha existente: {e}", file=sys.stderr)
        return set()


def aplicar_cabecalho(ws):
    """Formata a linha de cabeçalho."""
    fill = PatternFill("solid", fgColor=CABECALHO_COR)
    fonte = Font(bold=True, color=CABECALHO_FONTE_COR)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = fonte
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 22
    larguras = [50, 60, 20, 70, 20, 25]
    for i, (col, larg) in enumerate(zip(ws.iter_cols(min_row=1, max_row=1), larguras), 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = larg


def salvar_noticias(arquivo: str, noticias: list[dict], tema: str, links_existentes: set) -> int:
    """Adiciona notícias novas à planilha, retorna a quantidade inserida."""
    if os.path.exists(arquivo):
        wb = openpyxl.load_workbook(arquivo)
        ws = wb.active
        primeira_vez = False
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Notícias"
        ws.append(COLUNAS)
        aplicar_cabecalho(ws)
        primeira_vez = True

    agora = datetime.now().strftime("%Y-%m-%d %H:%M")
    inseridas = 0

    for n in noticias:
        link_norm = n["link"].strip()
        if link_norm not in links_existentes:
            ws.append([n["titulo"], link_norm, n["data"], n["resumo"], agora, tema])
            links_existentes.add(link_norm)
            inseridas += 1

    # Congela cabeçalho se for a primeira vez
    if primeira_vez:
        ws.freeze_panes = "A2"

    wb.save(arquivo)
    return inseridas


def main():
    args = parse_args()
    tema = args.tema.strip()
    arquivo = args.arquivo

    if not os.path.isabs(arquivo):
        arquivo = os.path.join(os.path.dirname(os.path.abspath(__file__)), arquivo)

    print(f"[{datetime.now():%Y-%m-%d %H:%M}] Buscando notícias sobre: '{tema}'")

    noticias = buscar_noticias_google(tema, args.max)
    print(f"  Encontradas: {len(noticias)} notícias")

    if not noticias:
        print("  Nenhuma notícia encontrada. Encerrando.")
        return

    links_existentes = carregar_links_existentes(arquivo)
    inseridas = salvar_noticias(arquivo, noticias, tema, links_existentes)

    print(f"  Novas inseridas: {inseridas} | Duplicatas ignoradas: {len(noticias) - inseridas}")
    print(f"  Arquivo salvo em: {arquivo}")


if __name__ == "__main__":
    main()
